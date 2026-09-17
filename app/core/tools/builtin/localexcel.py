# -*- coding: utf-8 -*-
import os
import json
import re
import difflib
from pathlib import Path
from typing import Any, List, Optional, Sequence
import openpyxl
from app.core.tools.base import BaseTool, ToolParameter


class _BaseExcelTool(BaseTool):
    """Excel 工具公共路径接地逻辑（不可注册，仅作复用基类）。"""

    def __init__(self, base_dir: Optional[str] = None) -> None:
        super().__init__()
        # 规范的数据根目录（通常是项目根）。注入后可用它解析相对路径，
        # 避免依赖进程 cwd；未注入时回退到本模块推断的仓库根。
        self._base_dir: Optional[Path] = Path(base_dir) if base_dir else None

    # ------------------------------------------------------------------
    # 路径接地：把相对路径稳定解析到规范根，且模型臆造路径时给出真实文件清单
    # ------------------------------------------------------------------
    def _canonical_base(self) -> Path:
        """本地 Excel 数据文件的规范根目录。"""
        if self._base_dir is not None:
            return self._base_dir
        # 本文件位于 app/core/tools/builtin/localexcel.py → parents[4] = 项目根
        return Path(__file__).resolve().parents[4]

    def _resolve_physical_path(self, file_path: str) -> Optional[Path]:
        """把模型传入的路径解析为真实存在的物理路径；找不到返回 None。

        ⚠️ 必须先剥掉开头的 ``/``：Windows 上 ``Path("/raw_data/x.xlsx").is_absolute()``
        是 **False**（只有盘符才算绝对），于是它走相对分支，而
        ``base / "/raw_data/x.xlsx"`` 在 Windows 语义下是**被右侧的 rooted 路径整体替换**
        成 ``D:/raw_data/x.xlsx`` → 必然找不到文件。模型很爱写前导斜杠（它会仿照
        file_list_tool 返回的 ``/raw_data/...`` 视图），实测为此白烧一整轮 LLM。
        """
        raw: str = str(file_path or "").strip().strip('"').strip("'")
        # rooted-but-driveless（/raw_data/... 或 \raw_data\...）统一按"相对项目根"处理
        if re.match(r"^[/\\](?![\\/])", raw) and not re.match(r"^[A-Za-z]:", raw):
            raw = raw.lstrip("/\\")
        if not raw:
            return None

        given = Path(raw)
        if given.is_absolute():
            return given if given.exists() else None
        base = self._canonical_base()
        candidates: List[Path] = [base / raw, Path(os.getcwd()) / raw]
        for candidate in candidates:
            if candidate.exists():
                return candidate
        return None

    def _grounding_with_real_files(self, file_path: str) -> str:
        """文件不存在时，返回真实可用的 Excel 文件清单，引导模型下一轮用真实路径。"""
        base = self._canonical_base()
        target_dir: Path = base / "raw_data" / "sales_intel"
        if not target_dir.exists():
            target_dir = base / "raw_data"
        xlsx_files: List[Path] = []
        if target_dir.exists():
            xlsx_files = sorted(
                p
                for p in target_dir.glob("**/*.xlsx")
                if not p.name.startswith("~$")  # 过滤 Excel 临时锁文件
            )
        if not xlsx_files:
            return (
                f"错误：找不到指定的 Excel 文件: {file_path}，"
                f"且当前未能定位到数据目录 {target_dir}，请先使用 file_list_tool / file_grep_tool 探查实际数据。"
            )

        lines: List[str] = [
            f"错误：找不到指定的 Excel 文件: {file_path}。"
            "以下是你当前真实可用的 Excel 文件（请严格按绝对路径使用，禁止臆造、改写或拼接文件名）：",
        ]
        wanted_name = Path(file_path).name
        best_match: Optional[Path] = max(
            xlsx_files,
            key=lambda fp: difflib.SequenceMatcher(None, wanted_name, fp.name).ratio(),
        )
        best_ratio = difflib.SequenceMatcher(None, wanted_name, best_match.name).ratio()
        if best_ratio >= 0.5:
            lines.append(f"- 与你要找的最相近的可能是: {best_match.resolve()}（相似度 {best_ratio:.2f}）")
        for fp in xlsx_files:
            rel = fp.relative_to(base).as_posix() if fp.is_relative_to(base) else str(fp)
            lines.append(f"\t绝对路径: {fp.resolve()} | 相对路径: {rel}")
        return "\n".join(lines)


def _preferred_excel_engine() -> Optional[str]:
    """优先用 calamine（Rust 实现，读大表快数倍）；未安装则返回 None 交给 pandas 自选。

    注意：``engine="calamine"`` 需要 ``python-calamine`` 包，未安装时**会直接抛错**，
    因此这里必须先探测再决定，不能写死（本机实测该包未安装）。
    """
    try:
        import python_calamine  # noqa: F401
        return "calamine"
    except Exception:  # noqa: BLE001 - 未安装 → 交给 pandas 默认引擎（xlsx→openpyxl）
        return None


# 名称接地阈值（与 _grounding_with_real_files 用的 0.5 同源；列名是精确标识，故略高）。
EXCEL_FUZZY_MIN_RATIO: float = 0.6
# 最优候选与次优候选的最小差距：差距太小说明"像的不止一个"，此时**绝不自动纠正**。
EXCEL_FUZZY_MIN_MARGIN: float = 0.10


class LocalExcelReadTool(_BaseExcelTool):
    """只读本地表格文件（pandas 统一读取；不触发人工审批）。"""

    name: str = "local_excel_read_tool"
    description: str = (
        "读取本地表格（xlsx/xls/csv/tsv/xlsm），返回结构摘要而不是整表内容："
        "不给 sheet_name 时列出**所有 sheet 的名称/行列数/列名/前若干行**；"
        "给 sheet_name 时看该表明细，可按列过滤（filter_column+filter_value）或直接做分组聚合"
        "（group_by+agg_column+agg_func，聚合在工具内完成，不要自己心算累加）；"
        "列名或 sheet 名写得不够准也会被自动就近纠正，并在结果里注明纠正了什么——"
        "**不要为了确认列名而反复调用本工具**，先按你记得的名字直接查，工具会告诉你真实列名。"
    )

    # 输出体积控制：摘要/明细都要给后续推理留余量
    DEFAULT_HEAD_ROWS: int = 5
    MAX_HEAD_ROWS: int = 50
    MAX_OUTPUT_CHARS: int = 6000

    def __init__(self, base_dir: Optional[str] = None) -> None:
        super().__init__(base_dir)
        self.parameters = [
            ToolParameter(name="file_path", type="string", description="本地表格文件的绝对路径或相对项目根的路径", required=True),
            ToolParameter(name="sheet_name", type="string",
                          description="工作表名；留空则返回该文件**所有 sheet 的结构摘要**", required=False),
            ToolParameter(name="cell", type="string",
                          description="单元格坐标（如 'D7'）；仅需单格取值时使用", required=False),
            ToolParameter(name="filter_column", type="string",
                          description="按列过滤行：列名（与 filter_value 搭配使用）", required=False),
            ToolParameter(name="filter_value", type="string",
                          description="过滤值：该列包含此文本的行才返回（不区分大小写）", required=False),
            ToolParameter(name="group_by", type="string",
                          description="分组聚合的分组列名（如 '产品线'）；配合 agg_column 使用", required=False),
            ToolParameter(name="agg_column", type="string",
                          description="聚合的目标列名（如 '金额(万元)'）", required=False),
            ToolParameter(name="agg_func", type="string",
                          description="聚合函数：sum（默认）/ mean / count / max / min / median", required=False),
            ToolParameter(name="head_rows", type="integer",
                          description=f"明细/摘要展示的行数，默认 {self.DEFAULT_HEAD_ROWS}，最大 {self.MAX_HEAD_ROWS}", required=False),
        ]

    # ------------------------------------------------------------------
    # 名称接地：把模型猜出来的列名/sheet 名收敛到真实存在的名字
    # ------------------------------------------------------------------
    @staticmethod
    def _ground_name(raw_name: str, candidates: Sequence[Any], role: str) -> "tuple[Any, str]":
        """把模型给出的名称接地到真实候选。返回 ``(真实候选 或 None, 说明文本)``。

        ⚠️ 为什么这是**必要护栏**而不是可选优化：
            planner 在**计划阶段**就要写出 ``tool_args_hint``（例如 group_by="产品线"），
            而表格的列名/sheet 名属于**运行时数据**——计划阶段它根本看不到，只能猜。
            猜错一次 = 工具报错 → 模型下一轮重试 → 白烧一整轮 LLM。
            而本轮的"直接复用 tool_args_hint"改动让这件事的风险**变大**了：hint 现在
            优先于带上下文（含上一步真实列名）的 FC 取参，于是 planner 的猜测会直达工具。
            与其回退那条优化，不如让工具自己把猜错接住。

        行为分三档，**核心是不瞎猜**：
            1. 精确命中 → 原样返回，不打扰模型；
            2. 唯一且足够相近 → 自动纠正，并**显式声明**"已将 X 解析为 Y"，让模型能自我校验；
            3. 相近候选不唯一 / 都不够像 → 返回 None + 真实候选清单，把决策权交回模型。

        注：``candidates`` 传原始对象（列名可能是数字、日期等非 str），只在比较与提示时
        字符串化，返回值仍是原始对象，保证 ``frame[resolved]`` 能正确索引。
        """
        labels: List[str] = [str(c) for c in candidates]
        if raw_name in labels:
            return candidates[labels.index(raw_name)], ""
        scored: List["tuple[str, float]"] = sorted(
            ((label, difflib.SequenceMatcher(None, raw_name, label).ratio()) for label in labels),
            key=lambda item: item[1],
            reverse=True,
        )
        if scored and scored[0][1] >= EXCEL_FUZZY_MIN_RATIO:
            best_label, best_ratio = scored[0]
            runner_up: float = scored[1][1] if len(scored) > 1 else 0.0
            if best_ratio - runner_up >= EXCEL_FUZZY_MIN_MARGIN:
                return candidates[labels.index(best_label)], (
                    f"{role} {raw_name!r} 未精确匹配，已按最相近的 {best_label!r} 解析"
                    f"（相似度 {best_ratio:.2f}）"
                )
        hints: str = "、".join(f"{label}（{ratio:.2f}）" for label, ratio in scored[:5])
        return None, (
            f"错误：{role} {raw_name!r} 不存在。实际可选: {', '.join(labels)}"
            + (f"\n最相近的候选: {hints}" if hints else "")
        )

    # ------------------------------------------------------------------
    # 读取 / 渲染
    # ------------------------------------------------------------------
    def _load_all_sheets(self, path: Path) -> "tuple[dict, str]":
        """读取整份表格 → ({sheet 名: DataFrame}, 引擎说明)。

        - csv/tsv 直接用 read_csv；
        - 其余走 pandas.read_excel(sheet_name=None) 一次拿到全部 sheet，**不需要猜 sheet 名**
          （旧实现默认 ``Sheet1`` + ``wb.active`` 兜底，模型猜错名字时静默返回错误的工作表）。
        """
        import pandas as pd

        suffix: str = path.suffix.lower()
        if suffix in (".csv", ".tsv"):
            sep: str = "\t" if suffix == ".tsv" else ","
            frame = pd.read_csv(path, sep=sep)
            return {path.stem: frame}, "pandas.read_csv"

        engine: Optional[str] = _preferred_excel_engine()
        read_kwargs: dict = {"sheet_name": None}
        if engine:
            read_kwargs["engine"] = engine
        sheets: dict = pd.read_excel(path, **read_kwargs)
        return sheets, f"pandas.read_excel(engine={engine or 'auto'})"

    @staticmethod
    def _fmt_cell(value: Any) -> str:
        import pandas as pd

        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ""
        text: str = str(value)
        return text if len(text) <= 60 else text[:60] + "…"

    def _render_frame_head(self, frame: Any, head_rows: int) -> str:
        """把 DataFrame 前 N 行渲染成紧凑的 `列|列` 表（比 JSON 省一半以上 token）。"""
        columns: List[str] = [str(c) for c in frame.columns]
        lines: List[str] = [" | ".join(columns)]
        for _, row in frame.head(head_rows).iterrows():
            lines.append(" | ".join(self._fmt_cell(row[c]) for c in frame.columns))
        return "\n".join(lines)

    def _preview_hint(self, sheet_name: str, rows: int, head_rows: int, *, overview: bool) -> str:
        """预览后的统一强提示（无条件附加，概览/明细两个分支共用，防止'预览没看到=数据不存在'误判）。"""
        shown = min(head_rows, rows)
        if rows > shown:
            prefix = f'再调用本工具并传 sheet_name="{sheet_name}" + ' if overview else "请用 "
            return (
                f"⚠️ 本表共 {rows} 行，仅展示前 {shown} 行；目标数据不在预览中**不代表不存在**，"
                f"严禁据此判定数据缺失。{prefix}filter_column+filter_value 精确筛选"
                "（filter_column 取上面列出的真实列名，filter_value 为该列的目标取值），"
                "统计请用 group_by+agg_column。"
            )
        return (
            f"本表共 {rows} 行，已全部展示；如需按条件取数仍可用 filter_column+filter_value，"
            "统计用 group_by+agg_column。"
        )

    def _render_overview(self, sheets: dict, head_rows: int) -> str:
        blocks: List[str] = []
        for name, frame in sheets.items():
            rows, cols = frame.shape
            block: str = (
                f"### {name}  ({rows} 行 × {cols} 列)\n"
                f"可用 filter_column（真实列名，任选其一作筛选列）: {', '.join(str(c) for c in frame.columns)}\n"
                f"前 {min(head_rows, rows)} 行:\n{self._render_frame_head(frame, head_rows)}\n"
                + self._preview_hint(name, rows, head_rows, overview=True)
            )
            blocks.append(block)
        return "\n\n".join(blocks)

    def _render_aggregate(self, frame: Any, group_by: str, agg_column: str, agg_func: str) -> str:
        """工具内完成分组聚合，避免把 50 行明细丢给模型心算（这是 token 与准确率的双重浪费）。"""
        import pandas as pd

        resolved_group, group_note = self._ground_name(group_by, list(frame.columns), "分组列")
        resolved_agg, agg_note = self._ground_name(agg_column, list(frame.columns), "聚合列")
        # 接不上就**不猜**：把真实列清单交回模型，让它下一轮用对的名字
        if resolved_group is None:
            return group_note
        if resolved_agg is None:
            return agg_note

        numeric = pd.to_numeric(frame[resolved_agg], errors="coerce")
        invalid_count: int = int(numeric.isna().sum() - frame[resolved_agg].isna().sum())
        working = frame.assign(__value__=numeric).dropna(subset=["__value__"])
        if working.empty:
            return f"错误：列 {resolved_agg!r} 没有可聚合的数值（可能全是文本/空值）。"

        grouped = working.groupby(resolved_group)["__value__"]
        if agg_func == "mean":
            series = grouped.mean()
        elif agg_func == "count":
            series = grouped.count()
        elif agg_func == "max":
            series = grouped.max()
        elif agg_func == "min":
            series = grouped.min()
        elif agg_func == "median":
            series = grouped.median()
        else:
            series = grouped.sum()
        ordered = series.sort_values(ascending=False)

        lines: List[str] = [
            # 接地说明放最前：模型看到"自己猜的名字被解析成了什么"才能自我校准
            *[f"（注：{note}）" for note in (group_note, agg_note) if note],
            f"## 聚合：按 {resolved_group!r} 对 {resolved_agg!r} 求 {agg_func}"
            f"（共 {len(ordered)} 组，降序）",
        ]
        for rank, (key, value) in enumerate(ordered.items(), 1):
            rendered = f"{value:,.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else str(value)
            lines.append(f"{rank}. {key}: {rendered}")
        if invalid_count > 0:
            lines.append(f"（提示：{resolved_agg!r} 中有 {invalid_count} 个非数值单元格未参与聚合）")
        return "\n".join(lines)

    def _render_filter(self, frame: Any, column: str, value: str, head_rows: int) -> str:
        resolved, note = self._ground_name(column, list(frame.columns), "过滤列")
        if resolved is None:
            return note
        prefix: str = f"（注：{note}）\n" if note else ""
        mask = frame[resolved].astype(str).str.contains(str(value), case=False, na=False, regex=False)
        hit = frame[mask]
        if hit.empty:
            return prefix + f"## 过滤：{resolved!r} 包含 {value!r} → 命中 0 行（该表共 {len(frame)} 行）"
        return (
            prefix
            + f"## 过滤：{resolved!r} 包含 {value!r} → 命中 {len(hit)} 行"
            + (f"（仅显示前 {head_rows} 行）" if len(hit) > head_rows else "")
            + f"\n{self._render_frame_head(hit, head_rows)}"
        )

    async def execute(self, **kwargs: Any) -> str:
        file_path = kwargs.get("file_path")
        sheet_name = str(kwargs.get("sheet_name") or "").strip()
        cell = str(kwargs.get("cell") or "").strip()
        filter_column = str(kwargs.get("filter_column") or "").strip()
        filter_value = str(kwargs.get("filter_value") or "")
        group_by = str(kwargs.get("group_by") or "").strip()
        agg_column = str(kwargs.get("agg_column") or "").strip()
        agg_func = str(kwargs.get("agg_func") or "sum").strip().lower()
        try:
            head_rows = max(1, min(self.MAX_HEAD_ROWS, int(kwargs.get("head_rows") or self.DEFAULT_HEAD_ROWS)))
        except (TypeError, ValueError):
            head_rows = self.DEFAULT_HEAD_ROWS

        if not file_path:
            return "错误：未提供文件路径 file_path"

        try:
            physical_path = self._resolve_physical_path(file_path)
            if physical_path is None:
                return self._grounding_with_real_files(file_path)

            # ── 情况 A：单格读取（保留 openpyxl：按坐标取值最精确，并顺带告知列名）──
            if cell:
                wb = openpyxl.load_workbook(str(physical_path), data_only=True)
                target_sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                value = target_sheet[cell].value
                column_letter: str = "".join(ch for ch in cell if ch.isalpha())
                column_name: Optional[Any] = None
                if target_sheet.max_row >= 1:
                    column_name = target_sheet.cell(row=1, column=openpyxl.utils.column_index_from_string(column_letter)).value
                return (
                    f"文件 [{os.path.basename(str(physical_path))}] 表 [{target_sheet.title}] "
                    f"单元格 {cell} = {value}"
                    + (f"（该列列名: {column_name}）" if column_name else "")
                )

            sheets, engine_note = self._load_all_sheets(physical_path)
            if not sheets:
                return f"错误：文件 [{physical_path.name}] 中没有任何可读工作表。"

            header: str = (
                f"文件: {physical_path.name} | {engine_note} | sheet 数: {len(sheets)}"
                f" | 可用 sheet: {', '.join(str(n) for n in sheets.keys())}"
            )

            # ── 情况 B：未指定 sheet → 全部 sheet 的结构摘要 ──
            # ⚠️ 例外：如果带了 filter/group 算子却没给 sheet，**绝不静默降级成概览**。
            # 实测事故（台账改负责人 trace）：模型传了 filter_column 但没传 sheet，
            # 旧逻辑直接忽略 filter 返回全表前 5 行，过滤白做、白烧一轮，还把"预览没命中"
            # 误判成"数据不存在"。单 sheet 文件自动选用；多 sheet 明确报错列清单。
            auto_sheet_note: str = ""
            if not sheet_name:
                if (filter_column and filter_value) or (group_by and agg_column):
                    if len(sheets) == 1:
                        sheet_name = str(next(iter(sheets)))
                        auto_sheet_note = f"未指定 sheet_name，该文件仅含一个 sheet [{sheet_name}]，已自动选用"
                    else:
                        return self._truncate(
                            header
                            + "\n\n错误：使用 filter_column+filter_value 或 group_by+agg_column "
                            "时必须提供 sheet_name（否则过滤/聚合参数会无法生效）。"
                            f"该文件有 {len(sheets)} 个 sheet："
                            f"{', '.join(str(n) for n in sheets.keys())}"
                        )
                else:
                    body = self._render_overview(sheets, head_rows)
                    return self._truncate(header + "\n\n" + body)

            # ── 情况 C：指定 sheet → 明细 / 过滤 / 聚合 ──
            sheet_note: str = auto_sheet_note
            if sheet_name not in sheets:
                # 容忍模型传了近名：唯一且够像才自动纠正，否则把真实 sheet 清单交回模型
                resolved_sheet, sheet_note = self._ground_name(
                    sheet_name, list(sheets.keys()), "sheet 名"
                )
                if resolved_sheet is None:
                    return sheet_note
                sheet_name = str(resolved_sheet)
            frame = sheets[sheet_name]
            parts: List[str] = [
                header,
                *([f"（注：{sheet_note}）"] if sheet_note else []),
                f"\n## Sheet [{sheet_name}]  {frame.shape[0]} 行 × {frame.shape[1]} 列",
                f"列: {', '.join(str(c) for c in frame.columns)}",
            ]

            if group_by and agg_column:
                parts.append("\n" + self._render_aggregate(frame, group_by, agg_column, agg_func))
            elif filter_column and filter_value:
                # 过滤已是用户主动收窄的结果：未显式给 head_rows 时一次展示到上限，
                # 避免"命中 10 行只给前 5 行 → 模型带 head_rows=50 再调一次"的重复往返。
                filter_rows = (
                    head_rows if kwargs.get("head_rows") not in (None, "")
                    else self.MAX_HEAD_ROWS
                )
                parts.append("\n" + self._render_filter(frame, filter_column, filter_value, filter_rows))
            else:
                parts.append(
                    f"\n前 {min(head_rows, frame.shape[0])} 行:\n{self._render_frame_head(frame, head_rows)}"
                )
                parts.append(
                    "\n" + self._preview_hint(sheet_name, int(frame.shape[0]), head_rows, overview=False)
                )
            return self._truncate("\n".join(parts))

        except Exception as e:
            return f"读取表格工具期间崩溃: {str(e)}"

    def _truncate(self, text: str) -> str:
        if len(text) <= self.MAX_OUTPUT_CHARS:
            return text
        return (
            text[: self.MAX_OUTPUT_CHARS]
            + f"\n（结果过长已截断，原长 {len(text)} 字符；请用 sheet_name/过滤/聚合收窄）"
        )


class LocalExcelWriteTool(_BaseExcelTool):
    """写本地表格的工具：单格更新 或 批量落表（写操作，在人工审批危险名单内）。"""

    name: str = "local_excel_write_tool"
    description: str = (
        "写入本地表格，三种模式："
        "(1)【推荐】语义定位更新 filter_column+filter_value+target_column+new_value："
        "按条件定位到唯一一行并修改某列（如把某客户的负责人改成某人），单元格坐标由工具内部计算，"
        "**不要自己算 A1/B3 这类坐标**（LLM 数列位置极易数错列，实测把 J 列数成 B 列写错字段）；"
        "(2) 单格更新 cell+value（仅在你已通过读取确认过坐标时使用）；"
        "(3) 批量写入 rows（JSON 数组），write_mode=append 追加 / replace 覆盖该 sheet。"
        "写操作需人工审批；语义模式要求 filter 条件必须唯一命中一行，命中 0 行或多行都会被拒绝。"
    )

    MAX_ROWS_PER_CALL: int = 5000

    def __init__(self, base_dir: Optional[str] = None) -> None:
        super().__init__(base_dir)
        self.parameters = [
            ToolParameter(name="file_path", type="string", description="目标表格文件路径（相对路径相对项目根解析）", required=True),
            ToolParameter(name="sheet_name", type="string", description="工作表名称；语义模式下多 sheet 文件必填，单 sheet 可省略", required=False),
            ToolParameter(name="filter_column", type="string",
                          description="语义更新模式：定位行的条件列名（如 '公司全称'）", required=False),
            ToolParameter(name="filter_value", type="string",
                          description="语义更新模式：定位行的条件值（该列包含此文本即命中，不区分大小写）；必须唯一命中一行",
                          required=False),
            ToolParameter(name="target_column", type="string",
                          description="语义更新模式：要修改的列名（如 '负责人'）", required=False),
            ToolParameter(name="new_value", type="string",
                          description="语义更新模式：写入的新值（数值列会自动按数字写入）", required=False),
            ToolParameter(name="cell", type="string",
                          description="单格模式：目标单元格坐标（如 'D7'）；语义模式无需提供", required=False),
            ToolParameter(name="value", type="string",
                          description="单格模式：写入该单元格的值", required=False),
            ToolParameter(name="rows", type="string",
                          description=(
                              "批量模式：JSON 数组。对象数组 [{\"列名\": 值, ...}]，"
                              "或二维数组 [[\"列名1\",\"列名2\"], [值, 值]]（二维时第一行必须是列名）。"
                              "一次可写多行，避免一格一格写入。"
                          ), required=False),
            ToolParameter(name="write_mode", type="string",
                          description="批量模式写入方式：append（默认，追加到 sheet 末尾）/ replace（清掉该 sheet 原有内容后重写）", required=False),
        ]

    # ------------------------------------------------------------------
    # 批量数据解析
    # ------------------------------------------------------------------
    @staticmethod
    def _parse_rows(raw: Any) -> "tuple[Optional[List[dict]], Optional[str]]":
        """把 rows 参数解析成 [{列名: 值}]；返回 (rows, 错误说明)。"""
        import ast

        if isinstance(raw, list):
            parsed: Any = raw
        else:
            text: str = str(raw or "").strip()
            if not text:
                return None, "rows 为空"
            parsed = None
            for parser in (json.loads, ast.literal_eval):
                try:
                    parsed = parser(text)
                    break
                except Exception:  # noqa: BLE001 - 逐个解析器试错
                    continue
            if parsed is None:
                return None, "rows 不是合法 JSON 数组（也尝试过 Python 字面量解析）"

        if isinstance(parsed, dict):
            parsed = [parsed]
        if not isinstance(parsed, list) or not parsed:
            return None, "rows 必须是非空数组"

        # 形态①：对象数组
        if all(isinstance(item, dict) for item in parsed):
            columns: List[str] = []
            for item in parsed:
                for key in item.keys():
                    if str(key) not in columns:
                        columns.append(str(key))
            return [{str(k): v for k, v in item.items()} for item in parsed], None

        # 形态②：二维数组（首行为列名）
        if all(isinstance(item, (list, tuple)) for item in parsed):
            header: List[str] = [str(c) for c in parsed[0]]
            body: List[dict] = []
            for row in parsed[1:]:
                body.append({header[i]: (row[i] if i < len(row) else None) for i in range(len(header))})
            if not body:
                return None, "二维数组只有表头没有数据行"
            return body, None

        return None, "rows 元素类型不支持（只接受对象数组或二维数组）"

    async def execute(self, **kwargs: Any) -> str:
        file_path = kwargs.get("file_path")
        # 语义模式需要区分"没给 sheet"（→ 单 sheet 自动选/多 sheet 报错）和"明确给了"；
        # cell/rows 老模式保持默认 Sheet1 的历史行为。
        raw_sheet = kwargs.get("sheet_name")
        sheet_name = str(raw_sheet or "Sheet1")
        cell = kwargs.get("cell")
        value = kwargs.get("value")
        rows_raw = kwargs.get("rows")
        write_mode = str(kwargs.get("write_mode") or "append").strip().lower()
        if write_mode not in ("append", "replace"):
            write_mode = "append"
        filter_column = str(kwargs.get("filter_column") or "").strip()
        filter_value = kwargs.get("filter_value")
        target_column = str(kwargs.get("target_column") or "").strip()
        new_value = kwargs.get("new_value")

        if not file_path:
            return "错误：未提供文件路径 file_path"

        try:
            target = Path(file_path)
            if not target.is_absolute():
                target = self._canonical_base() / file_path
            target_path = str(target)

            # ── 模式①：批量写入 ──
            if rows_raw not in (None, "", [], "[]"):
                return self._bulk_write(target_path, sheet_name, rows_raw, write_mode)

            # ── 模式②：语义定位更新（条件定位行+列名定位列，坐标工具内部算）──
            # 这是修改"某条记录某字段"的推荐路径：让 LLM 算 A1 坐标不可靠
            # （实测把第 10 列"负责人"数成 B 列，把公司名写成了人名）。
            if filter_column or target_column or filter_value not in (None, "") or new_value is not None:
                if not (filter_column and filter_value not in (None, "") and target_column):
                    return (
                        "错误：语义更新模式必须同时提供 filter_column + filter_value + target_column"
                        "（+ new_value）四个参数；只改单个已知坐标的格子请用 cell+value 模式"
                    )
                return self._semantic_update(
                    target_path,
                    str(raw_sheet).strip() if raw_sheet not in (None, "") else "",
                    filter_column,
                    filter_value,
                    target_column,
                    new_value,
                )

            # ── 模式③：单格写入（原有语义，审批后立即落盘）──
            if not cell:
                return (
                    "错误：未提供任何写入参数。修改某条记录的字段请用 "
                    "filter_column+filter_value+target_column+new_value 语义模式；"
                    "已知坐标改单格才用 cell+value；一次写多行用 rows"
                )
            if os.path.exists(target_path):
                wb = openpyxl.load_workbook(target_path)
            else:
                wb = openpyxl.Workbook()
            sheet = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(title=sheet_name)
            sheet[cell] = value
            wb.save(target_path)
            return (
                f"成功：已在文件 [{os.path.basename(target_path)}] 的 [{sheet.title}] 表 "
                f"{cell} 位置写入值: {value}"
            )

        except Exception as e:
            return f"写入 Excel 工具期间崩溃: {str(e)}"

    def _semantic_update(
        self,
        target_path: str,
        sheet_name: str,
        filter_column: str,
        filter_value: Any,
        target_column: str,
        new_value: Any,
    ) -> str:
        """按"条件列=值"定位唯一行、按列名定位目标列，工具内部计算 A1 坐标后写入。

        为什么坐标必须在工具内算而不是让 LLM 传：
            trace 实测模型把 15 列里的第 10 列"负责人"数成了 B 列（第 2 列"公司全称"），
            一次写错字段。列名→列字母是确定性映射，放进工具里做既不需要 LLM、也不会错。
        写操作安全约束：
            - 条件必须**唯一命中一行**：0 命中 → 拒绝（防条件写错静默失败）；
              多命中 → 拒绝并列出行号（防一次误改一批）；
            - pandas 负责定位（与读工具同一套接地/contains 语义），
              openpyxl 负责落盘（保留原表格式/公式）；
            - pandas 行序即 Excel 数据行序（第 1 行表头），位置 i → Excel 行 i+2。
        """
        import pandas as pd

        if not os.path.exists(target_path):
            return f"错误：语义更新只能修改已存在的文件，目标文件不存在：{target_path}"
        suffix: str = Path(target_path).suffix.lower()
        if suffix not in (".xlsx", ".xlsm"):
            return (
                f"错误：语义更新目前仅支持 .xlsx/.xlsm（当前 {suffix or '无后缀'}）；"
                "老格式 .xls 请先另存为 .xlsx"
            )

        # 1) 读全量数据 + sheet 选择（复用读工具的加载/接地实现，保证读写语义一致）
        reader = LocalExcelReadTool(str(self._canonical_base()))
        sheets, _engine_note = reader._load_all_sheets(Path(target_path))
        if not sheets:
            return f"错误：文件 [{os.path.basename(target_path)}] 中没有任何可读工作表。"
        if not sheet_name:
            if len(sheets) == 1:
                sheet_name = str(next(iter(sheets)))
            else:
                return (
                    f"错误：该文件有 {len(sheets)} 个 sheet，语义更新必须提供 sheet_name："
                    f"{', '.join(str(n) for n in sheets.keys())}"
                )
        if sheet_name not in sheets:
            resolved, note = LocalExcelReadTool._ground_name(
                sheet_name, list(sheets.keys()), "sheet 名"
            )
            if resolved is None:
                return note
            sheet_name = str(resolved)
        frame = sheets[sheet_name]

        # 2) 条件列 / 目标列接地到真实列名
        resolved_filter, filter_note = LocalExcelReadTool._ground_name(
            filter_column, list(frame.columns), "条件列"
        )
        if resolved_filter is None:
            return filter_note
        resolved_target, target_note = LocalExcelReadTool._ground_name(
            target_column, list(frame.columns), "目标列"
        )
        if resolved_target is None:
            return target_note

        # 3) 定位行：与读工具 _render_filter 同一匹配语义（包含、不区分大小写）
        mask = frame[resolved_filter].astype(str).str.contains(
            str(filter_value), case=False, na=False, regex=False
        )
        hit_positions: List[int] = [i for i, hit in enumerate(mask) if bool(hit)]
        if not hit_positions:
            return (
                f"错误：条件 {resolved_filter!r} 包含 {str(filter_value)!r} 在表 [{sheet_name}] "
                f"中命中 0 行（该表共 {len(frame)} 行），未做任何修改。"
                f"请先用 local_excel_read_tool 核对 {resolved_filter!r} 列的真实取值。"
            )
        if len(hit_positions) > 1:
            sample = [
                f"Excel 第 {pos + 2} 行（{resolved_filter}={frame[resolved_filter].iloc[pos]!r}）"
                for pos in hit_positions[:5]
            ]
            return (
                f"错误：条件 {resolved_filter!r} 包含 {str(filter_value)!r} 命中了 {len(hit_positions)} 行，"
                "写操作要求唯一命中，已拒绝修改（防止误改多条记录）。请收紧 filter_value 后重试。"
                f"命中行：{'; '.join(sample)}"
            )
        data_row_pos: int = hit_positions[0]
        excel_row: int = data_row_pos + 2  # 第 1 行是表头

        # 4) 打开工作簿定位列序号与单元格（列名以表头第 1 行为准，逐列比对真实列名）
        wb = openpyxl.load_workbook(target_path)
        if sheet_name not in wb.sheetnames:
            return f"错误：pandas 可读到 sheet {sheet_name!r}，但 openpyxl 打开后未找到，未做修改。"
        ws = wb[sheet_name]
        excel_col: Optional[int] = None
        target_label = str(resolved_target)
        for col_idx in range(1, ws.max_column + 1):
            if str(ws.cell(row=1, column=col_idx).value) == target_label:
                excel_col = col_idx
                break
        if excel_col is None:
            return f"错误：在工作表第 1 行表头中找不到列 {target_label!r}，未做修改。"

        cell_ref: str = f"{openpyxl.utils.get_column_letter(excel_col)}{excel_row}"
        old_value = ws[cell_ref].value

        # 5) 类型适配：目标列是数值列时，尽量把入参写成数字而不是字符串
        write_value: Any = new_value
        if new_value is not None and pd.api.types.is_numeric_dtype(frame[resolved_target]):
            text = str(new_value).strip().replace(",", "")
            try:
                num = float(text)
                write_value = int(num) if num.is_integer() else num
            except (TypeError, ValueError):
                write_value = new_value  # 转不了就按字符串写，并在回执里提示

        ws[cell_ref] = write_value
        wb.save(target_path)

        notes: List[str] = [n for n in (filter_note, target_note) if n]
        return (
            f"成功：[{os.path.basename(target_path)}] 表 [{sheet_name}] 中 "
            f"{resolved_filter}={frame[resolved_filter].iloc[data_row_pos]!r} 的唯一一行"
            f"（Excel 第 {excel_row} 行），列 {resolved_target!r}（单元格 {cell_ref}）"
            f"已由 {old_value!r} 更新为 {write_value!r}。"
            + ("".join(f"（注：{n}）" for n in notes) if notes else "")
        )

    def _bulk_write(self, target_path: str, sheet_name: str, rows_raw: Any, write_mode: str) -> str:
        """批量落表。

        两条实现路径，按语义分开——这不是重复造轮子，而是"不要为了快把台账写坏"：

        - **append 到已存在的 sheet**：用 openpyxl ``ws.append()`` 逐行追加，
          保留原有单元格格式/公式/列宽（张台账往往有格式约定，整表重写会全部丢掉）；
        - **replace / 新建 sheet**：用 ``pandas.ExcelWriter``（engine=openpyxl）整表写出，
          比一格一格 openpyxl 赋值快得多，且列对齐由 pandas 保证。

        ⚠️ 引擎选择说明：pandas 的 xlsxwriter 引擎**只能新建文件**，无法写入/修改既有工作簿
        （``mode="a"`` 不被 xlsxwriter 支持），因此这里的 ExcelWriter 固定用 openpyxl 引擎。
        """
        import pandas as pd

        records, parse_error = self._parse_rows(rows_raw)
        if records is None:
            return f"错误：{parse_error}"
        if len(records) > self.MAX_ROWS_PER_CALL:
            return f"错误：单次最多写入 {self.MAX_ROWS_PER_CALL} 行（本次 {len(records)} 行），请分批。"

        columns: List[str] = []
        for record in records:
            for key in record.keys():
                if key not in columns:
                    columns.append(key)
        frame = pd.DataFrame(records, columns=columns)

        file_exists: bool = os.path.exists(target_path)
        sheet_exists: bool = False
        if file_exists:
            try:
                with pd.ExcelFile(target_path, engine="openpyxl") as handle:
                    sheet_exists = sheet_name in handle.sheet_names
            except Exception:  # noqa: BLE001 - 读不出来就按"不存在"处理，走新建分支
                sheet_exists = False

        if write_mode == "append" and sheet_exists:
            wb = openpyxl.load_workbook(target_path)
            sheet = wb[sheet_name]
            for record in records:
                sheet.append([record.get(column) for column in columns])
            wb.save(target_path)
            return (
                f"成功（追加）：[{os.path.basename(target_path)}] 的 [{sheet_name}] 表末尾追加了 "
                f"{len(records)} 行，列: {', '.join(columns)}（原有格式与内容保留）"
            )

        with pd.ExcelWriter(
            target_path,
            engine="openpyxl",
            mode="a" if file_exists else "w",
            if_sheet_exists="replace" if file_exists else None,
        ) as writer:
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        return (
            f"成功（{'覆盖' if write_mode == 'replace' else '新建'}）："
            f"[{os.path.basename(target_path)}] 的 [{sheet_name}] 表写入了 {len(records)} 行 × "
            f"{len(columns)} 列，列: {', '.join(columns)}"
        )
